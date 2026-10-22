"""Поиск по файлам по смыслу (аналог семантического поиска Windows).

  python -m winai files index "C:\\Users\\me\\Documents" [--images] [--watch 30]
  python -m winai files search "договор аренды прошлым летом"

Индекс инкрементальный: изменённые файлы переиндексируются, удалённые — убираются.
"""
import os
import sqlite3
import time
from pathlib import Path

import numpy as np

from . import config, power, vecsearch

DB = config.DATA / "files.db"
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".ini", ".toml", ".yaml", ".yml", ".xml",
            ".html", ".htm", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go",
            ".rs", ".sql", ".bat", ".ps1", ".sh", ".rtf", ".tex"}
DOC_EXT = {".pdf", ".docx", ".pptx", ".xlsx"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", "AppData", "$Recycle.Bin",
             "site-packages", ".cache", "models"}
MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 60_000


def db():
    DB.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB)
    con.executescript("""
    CREATE TABLE IF NOT EXISTS files(id INTEGER PRIMARY KEY, path TEXT UNIQUE, mtime REAL, size INTEGER);
    CREATE TABLE IF NOT EXISTS chunks(id INTEGER PRIMARY KEY, file_id INTEGER, text TEXT, vec BLOB);
    CREATE INDEX IF NOT EXISTS chunks_file ON chunks(file_id);
    CREATE VIRTUAL TABLE IF NOT EXISTS cfts USING fts5(text, content='chunks', content_rowid='id');
    """)
    return con


def extract(path: Path, images=False) -> str:
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        raw = path.read_bytes()[:MAX_CHARS * 4]
        for enc in ("utf-8", "utf-16", "cp1251"):
            try:
                return raw.decode(enc)[:MAX_CHARS]
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", "ignore")[:MAX_CHARS]
    if ext == ".pdf":
        from pypdf import PdfReader
        out = []
        for page in PdfReader(str(path)).pages[:60]:
            out.append(page.extract_text() or "")
            if sum(map(len, out)) > MAX_CHARS:
                break
        return "\n".join(out)[:MAX_CHARS]
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)[:MAX_CHARS]
    if ext == ".pptx":
        from pptx import Presentation
        parts = [sh.text_frame.text for s in Presentation(str(path)).slides
                 for sh in s.shapes if sh.has_text_frame]
        return "\n".join(parts)[:MAX_CHARS]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=500, values_only=True):
                parts.append(" ".join(str(v) for v in row if v is not None))
        return "\n".join(parts)[:MAX_CHARS]
    if images and ext in IMG_EXT:
        from PIL import Image

        from .ocr import ocr
        return ocr(Image.open(path))
    return ""


def walk(roots, images=False):
    exts = TEXT_EXT | DOC_EXT | (IMG_EXT if images else set())
    for root in roots:
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
            for f in filenames:
                p = Path(dirpath) / f
                if p.suffix.lower() in exts:
                    yield p


def index(roots, images=False, emb=None, log=print):
    from .embedder import shared
    emb = emb or shared()
    con = db()
    seen, added = set(), 0
    roots = [str(Path(r).resolve()) for r in roots]
    for p in walk(roots, images):
        try:
            st = p.stat()
        except OSError:
            continue
        key = str(p)
        seen.add(key)
        if st.st_size > MAX_BYTES:
            continue
        row = con.execute("SELECT id, mtime FROM files WHERE path=?", (key,)).fetchone()
        if row and abs(row[1] - st.st_mtime) < 1:
            continue
        try:
            text = extract(p, images)
        except Exception as e:
            log(f"пропуск {p.name}: {type(e).__name__}")
            continue
        if row:
            _forget(con, row[0])
        fid = con.execute("INSERT OR REPLACE INTO files(path,mtime,size) VALUES(?,?,?)",
                          (key, st.st_mtime, st.st_size)).lastrowid
        # имя файла тоже ищется по смыслу
        pieces = emb.chunk(f"{p.stem.replace('_', ' ')}. {text}", max_chunks=12, words=90)
        for piece, vec in zip(pieces, emb.embed_passages(pieces)):
            cid = con.execute("INSERT INTO chunks(file_id,text,vec) VALUES(?,?,?)",
                              (fid, piece, vec.astype(np.float32).tobytes())).lastrowid
            con.execute("INSERT INTO cfts(rowid,text) VALUES(?,?)", (cid, piece))
        added += 1
        if added % 20 == 0:
            con.commit()
            log(f"проиндексировано {added}...")
            if power.on_battery():
                time.sleep(0.5)  # от батареи не гоним NPU на максимум
    # убрать файлы, которых больше нет
    for fid, path in con.execute("SELECT id, path FROM files").fetchall():
        if any(path.startswith(r) for r in roots) and path not in seen:
            _forget(con, fid)
            con.execute("DELETE FROM files WHERE id=?", (fid,))
    con.commit()
    log(f"готово: новых/изменённых файлов {added}, всего {con.execute('SELECT count(*) FROM files').fetchone()[0]}")
    return added


def _forget(con, fid):
    for cid, text in con.execute("SELECT id, text FROM chunks WHERE file_id=?", (fid,)).fetchall():
        con.execute("INSERT INTO cfts(cfts, rowid, text) VALUES('delete', ?, ?)", (cid, text))
    con.execute("DELETE FROM chunks WHERE file_id=?", (fid,))


def search(query: str, k: int = 15, emb=None):
    from .embedder import shared
    emb = emb or shared()
    con = db()
    top = vecsearch.hybrid(con, emb.embed_query(query), query, "SELECT id, vec FROM chunks", "cfts", k * 4)
    best = {}
    for cid, score in top:
        r = con.execute("SELECT f.path, c.text FROM chunks c JOIN files f ON f.id=c.file_id WHERE c.id=?",
                        (cid,)).fetchone()
        if r and (r[0] not in best or best[r[0]]["score"] < score):
            best[r[0]] = {"path": r[0], "score": score, "snippet": r[1][:300]}
    return sorted(best.values(), key=lambda x: -x["score"])[:k]


def main(argv):
    import argparse
    ap = argparse.ArgumentParser(prog="winai files")
    sub = ap.add_subparsers(dest="cmd", required=True)
    i = sub.add_parser("index")
    i.add_argument("folders", nargs="+")
    i.add_argument("--images", action="store_true", help="распознавать текст на картинках (медленнее)")
    i.add_argument("--watch", type=float, default=0, help="переиндексировать каждые N минут")
    s = sub.add_parser("search")
    s.add_argument("query")
    s.add_argument("-k", type=int, default=10)
    a = ap.parse_args(argv)
    if a.cmd == "index":
        while True:
            index(a.folders, a.images)
            if not a.watch:
                break
            time.sleep(a.watch * 60)
    else:
        for r in search(a.query, a.k):
            print(f"{r['score']:.2f} {r['path']}\n      {r['snippet'][:120]!r}")
